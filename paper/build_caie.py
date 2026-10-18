"""Builds the Computers & Industrial Engineering submission package into ../CAIE_submission."""
import json, re, shutil, subprocess, pathlib
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
P = pathlib.Path(__file__).parent
OUT = P.parent / "CAIE_submission"; OUT.mkdir(exist_ok=True)
L = json.loads((P / "links.json").read_text())
def fill(t):
    for k, v in L.items(): t = t.replace(k, v if v else "[%s to be added]" % k.replace("_", " ").lower())
    return t
# ---------- manuscript (anonymised: no author block) ----------
t = (P / "manuscript_src.md").read_text()
t = re.sub(r'\{\{(\w+)\}\}', lambda m: (P / "tables" / f"{m.group(1)}.md").read_text().strip(), t)
(P / "ms_caie.md").write_text(t)
subprocess.run(["pandoc", "ms_caie.md", "-o", "ms_caie.docx", "--resource-path=."], cwd=P, check=True)
d = Document(str(P / "ms_caie.docx"))
for s in d.sections:
    s.left_margin = s.right_margin = Cm(2.5); s.top_margin = s.bottom_margin = Cm(2.5)
    ln = OxmlElement("w:lnNumType"); ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous"); s._sectPr.append(ln)
def setfont(st, size=None, bold=None):
    st.font.name = "Times New Roman"; st.font.color.rgb = RGBColor(0, 0, 0)
    if st.element.rPr is not None: st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if size: st.font.size = Pt(size)
    if bold is not None: st.font.bold = bold
for n in ["Normal", "Body Text", "First Paragraph", "Compact", "Table Caption", "Image Caption", "Captioned Figure"]:
    try:
        st = d.styles[n]; setfont(st, size=12 if n != "Compact" else 9)
        if n in ("Body Text", "First Paragraph"): st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(6)
    except KeyError: pass
for n, sz in [("Title", 16), ("Heading 1", 14), ("Heading 2", 13), ("Heading 3", 12)]:
    try: setfont(d.styles[n], size=sz, bold=True)
    except KeyError: pass
for p in d.paragraphs:
    if p.style.name == "Title": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for tbl in d.tables:
    tp = tbl._tbl.tblPr; bd = OxmlElement("w:tblBorders")
    for e, v in [("top", 8), ("bottom", 8), ("insideH", 2)]:
        x = OxmlElement(f"w:{e}"); x.set(qn("w:val"), "single"); x.set(qn("w:sz"), str(v)); x.set(qn("w:color"), "000000"); bd.append(x)
    tp.append(bd); nc = len(tbl.columns); fs = 8 if nc >= 7 else (9 if nc >= 5 else 10)
    for ri, row in enumerate(tbl.rows):
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1); p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    r.font.size = Pt(fs); r.font.name = "Times New Roman"
                    if ri == 0: r.font.bold = True
d.core_properties.author = ""; d.core_properties.last_modified_by = ""; d.core_properties.title = "Manuscript (anonymised)"
d.save(str(OUT / "03_Manuscript_anonymised.docx"))
subprocess.run(["pandoc", "ms_caie.md", "-s", "--mathml", "--embed-resources", "--resource-path=.", "-c", "pdf.css", "-o", "ms_caie.html", "--metadata", "pagetitle=Manuscript"], cwd=P, check=True)
from playwright.sync_api import sync_playwright
with sync_playwright() as pw:
    b = pw.chromium.launch(executable_path="/opt/pw-browsers/chromium"); pg = b.new_page()
    pg.goto("file://" + str(P / "ms_caie.html")); pg.wait_for_timeout(1500)
    pg.pdf(path=str(OUT / "03_Manuscript_anonymised.pdf"), format="A4", print_background=True, prefer_css_page_size=True,
           display_header_footer=True, header_template="<span></span>",
           footer_template='<div style="font-size:8px;width:100%;text-align:center"><span class="pageNumber"></span></div>')
    b.close()
# ---------- helpers for small docs ----------
def newdoc(size=12):
    d = Document()
    for s in d.sections: s.left_margin = s.right_margin = Cm(2.5); s.top_margin = s.bottom_margin = Cm(2.3)
    st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(size)
    return d
def para(d, text, bold=False, after=8, hl=False, italic=False):
    p = d.add_paragraph(); r = p.add_run(text); r.bold = bold; r.italic = italic
    if hl: r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    p.paragraph_format.space_after = Pt(after); return p
TITLE = ("Beyond Forecast Accuracy: Decision-Centric Evaluation of Demand Forecasting for Inventory "
         "Replenishment under Intermittent and Shifting Retail Demand")
# ---------- highlights ----------
H = ["Accuracy and inventory cost rank demand forecasters differently",
     "Accuracy-picked forecaster with a textbook rule costs 140% more",
     "A fully tuned normal-theory rule is still 26% above the optimum",
     "Pooled empirical calibration raised cost; trained quantiles cut it",
     "Per-SKU strategy selection backfires against one portfolio choice"]
assert all(len(h) <= 85 for h in H), [len(h) for h in H]
d = newdoc(); para(d, "Highlights", True, 10)
for h in H:
    p = d.add_paragraph(h, style="List Bullet"); p.paragraph_format.space_after = Pt(4)
d.save(str(OUT / "02_Highlights.docx"))
# ---------- title page ----------
d = newdoc(); para(d, TITLE, True, 12)
p = para(d, "", after=2); r = p.add_run("Md Ibtashamur Rahman"); r.bold = True; r2 = p.add_run("a"); r2.font.superscript = True; r2.bold = True
p = para(d, "", after=8); r = p.add_run("a"); r.font.superscript = True
p.add_run("Department of Industrial and Production Engineering, Shahjalal University of Science and Technology, Kumargaon, Sylhet 3114, Bangladesh")
para(d, "Corresponding author: Md Ibtashamur Rahman", after=2)
para(d, "E-mail: ibtarahman@gmail.com", after=2)
para(d, "ORCID: https://orcid.org/0009-0008-8641-9289", after=10)
para(d, "Acknowledgments", True, 2); para(d, "None.", after=8)
para(d, "Author contribution (CRediT)", True, 2)
para(d, "Md Ibtashamur Rahman: Conceptualization, Methodology, Software, Validation, Formal analysis, Investigation, Data curation, Writing – original draft, Writing – review & editing, Visualization.", after=8)
para(d, "Declaration of competing interest", True, 2)
para(d, "The author declares that he has no known competing financial interests or personal relationships that could have appeared to influence the work reported in this paper.", after=8, hl=True)
para(d, "Funding", True, 2)
para(d, "This research did not receive any specific grant from funding agencies in the public, commercial, or not-for-profit sectors.", after=8, hl=True)
para(d, "Declaration of generative AI and AI-assisted technologies in the writing process", True, 2)
para(d, "During the preparation of this work the author used Claude (Anthropic) to assist with code development, analysis scripts, and the drafting and editing of the manuscript. After using this tool, the author reviewed and edited the content as needed and takes full responsibility for the content of the published article. (This statement also appears in the manuscript before the references.)", after=8, hl=True)
para(d, "Data and code availability", True, 2)
para(d, "The M5 data are public. Code, configuration files, result tables and the reproduction script: https://github.com/Ibtasham/M5-decision-centric-forecasting, archived at https://doi.org/10.5281/zenodo.23450902 (Rahman, 2026). In the anonymised manuscript this reference appears as Anonymous (2026).", after=8)
para(d, "Software reference: Rahman, M. I. (2026). Code and results for the study: decision-centric evaluation of demand forecasting [Computer software]. Zenodo. https://doi.org/10.5281/zenodo.23450902", after=8)
para(d, "Yellow items are drafts that only the author can confirm. Remove the highlight once checked.", italic=True)
d.save(str(OUT / "01_Title_page.docx"))
# ---------- cover letter ----------
d = newdoc(11.5)
para(d, "18 October 2026", after=14); para(d, "The Editor-in-Chief", after=2); para(d, "Computers & Industrial Engineering", after=14)
para(d, "Dear Editor,")
para(d, f"I submit my manuscript “{TITLE}” for consideration as a research article in Computers & Industrial Engineering.")
para(d, "Organisations invest heavily in forecast accuracy, but the value of a forecast comes through inventory decisions. Using 3,999 real retail series from the public M5 data and a lost-sales inventory simulation, I show that the forecaster selected on out-of-sample accuracy and run through the standard normal-theory base-stock rule costs 140.4% more than the cost-optimal forecaster–policy pairing (95% CI [128.8, 153.4]), at a lower fill rate. Replacing the forecaster inside that rule is worth 0.4% on the test block; replacing the rule is worth 58.4%.")
para(d, "The paper also reports an honest tuned baseline (a validation-tuned normal-theory rule is still 25.8% above the optimum), shows that a pooled empirical calibration raises cost while trained quantiles lower it, and shows that per-SKU strategy selection is counter-productive. The findings hold across ten decision environments and on an independent sample draw.")
para(d, "I believe the paper suits Computers & Industrial Engineering because it addresses an industrial-engineering decision problem, inventory replenishment under intermittent demand, with a simulation-based evaluation and practical guidance for selecting forecasting and replenishment methods. Every number is produced by publicly available code.")
para(d, fill("Code and results: REPO_URL (archived at ZENODO_DOI)."))
para(d, "The manuscript is original, has not been published, and is not under consideration elsewhere. I am the sole author and approve the submission.")
para(d, "Thank you for your consideration.", after=14)
para(d, "Sincerely,", after=2); para(d, "Md Ibtashamur Rahman", True, 0)
para(d, "Department of Industrial and Production Engineering", after=0)
para(d, "Shahjalal University of Science and Technology, Sylhet 3114, Bangladesh", after=0)
d.save(str(OUT / "04_Cover_letter.docx"))
# ---------- figures ----------
F = OUT / "Figures"; F.mkdir(exist_ok=True); R = P.parent / "results"
for n, src in [(1, "fig0_pipeline.png"), (2, "main/fig1_leverage.png"), (3, "main/fig6_accuracy_scatter.png"), (4, "main/fig2_decomposition.png"),
               (5, "main/fig4_selection.png"), (6, "fig5_sensitivity.png"), (7, "main/fig3_frontier.png")]:
    shutil.copy(R / src, F / f"Figure_{n}.png")
subprocess.run(["pandoc", "EDITOR_DIAGNOSIS_AND_DEFENCE_GUIDE.md", "-o", str(OUT / "Author_only_Diagnosis_and_Defence_Guide.docx")], cwd=P, check=True)
print("built", sorted(x.name for x in OUT.iterdir()))
