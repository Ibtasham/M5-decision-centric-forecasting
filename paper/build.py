import re,subprocess,pathlib
P=pathlib.Path(__file__).parent
t=(P/'manuscript_src.md').read_text()
def sub(m): return (P/'tables'/f'{m.group(1)}.md').read_text().strip()
t=re.sub(r'\{\{(\w+)\}\}',sub,t)
import json
_L=json.loads((P/'links.json').read_text())
for k,v in _L.items():
    t=t.replace(k, v if v else '[%s to be added]'%k.replace('_',' ').lower())
(P/'manuscript.md').write_text(t)

# ---- docx
subprocess.run(['pandoc','manuscript.md','-o','manuscript_submission.docx','--resource-path=.'],cwd=P,check=True)
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
d=Document(str(P/'manuscript_submission.docx'))
for s in d.sections:
    s.left_margin=s.right_margin=Cm(2.2); s.top_margin=s.bottom_margin=Cm(2.2)
def setfont(st,name='Times New Roman',size=None,bold=None,color=True):
    st.font.name=name
    st.element.rPr.rFonts.set(qn('w:eastAsia'),name) if st.element.rPr is not None else None
    if size: st.font.size=Pt(size)
    if bold is not None: st.font.bold=bold
    if color: st.font.color.rgb=RGBColor(0,0,0)
for n in ['Normal','Body Text','First Paragraph','Compact','Table Caption','Image Caption','Captioned Figure']:
    try:
        st=d.styles[n]; setfont(st,size=11 if n!='Compact' else 9)
        if n in('Body Text','First Paragraph'):
            st.paragraph_format.space_after=Pt(6); st.paragraph_format.line_spacing=1.15
    except KeyError: pass
for n,sz in [('Title',17),('Heading 1',14),('Heading 2',13),('Heading 3',11.5),('Author',12)]:
    try: setfont(d.styles[n],size=sz,bold=True)
    except KeyError: pass
try: d.styles['Image Caption'].font.size=Pt(9.5); d.styles['Image Caption'].font.italic=True
except KeyError: pass
for p in d.paragraphs:
    if p.style.name in('Title','Author'): p.alignment=WD_ALIGN_PARAGRAPH.CENTER
def borders(tbl):
    tblPr=tbl._tbl.tblPr
    b=OxmlElement('w:tblBorders')
    for e,v in [('top',8),('bottom',8),('insideH',2)]:
        x=OxmlElement(f'w:{e}'); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),str(v)); x.set(qn('w:color'),'000000'); b.append(x)
    tblPr.append(b)
for tbl in d.tables:
    borders(tbl)
    ncol=len(tbl.columns); fs=8 if ncol>=7 else (9 if ncol>=5 else 10)
    for ri,row in enumerate(tbl.rows):
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(1); p.paragraph_format.space_before=Pt(1)
                for r in p.runs:
                    r.font.size=Pt(fs); r.font.name='Times New Roman'
                    if ri==0: r.font.bold=True
d.save(str(P/'manuscript_submission.docx'))
subprocess.run(['pandoc','manuscript.md','-s','--mathml','--embed-resources','--resource-path=.','-c','pdf.css','-o','manuscript.html','--metadata','pagetitle=Manuscript'],cwd=P,check=True)
from playwright.sync_api import sync_playwright
with sync_playwright() as pw:
    b=pw.chromium.launch(executable_path='/opt/pw-browsers/chromium'); pg=b.new_page()
    pg.goto('file://'+str(P/'manuscript.html')); pg.wait_for_timeout(1500)
    pg.pdf(path=str(P/'manuscript_submission.pdf'),format='A4',print_background=True,prefer_css_page_size=True,
      display_header_footer=True,header_template='<span></span>',
      footer_template='<div style="font-size:8px;width:100%;text-align:center"><span class="pageNumber"></span></div>')
    b.close()
